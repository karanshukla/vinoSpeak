"""Writes the document fixtures. The tests don't run it: they read each
NAME.txt, what vinoSearch's Control1.FileText gave for NAME, and compare
`vinospeak --print-text --kind document` on it with NAME.expected. After a
change here, index the new files with vinoSearch and redo both.

uv run --no-project --with python-docx --with python-pptx --with reportlab tests/fixtures/documents/make.py
"""

from pathlib import Path

from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import A5
from reportlab.pdfgen.canvas import Canvas

HERE = Path(__file__).parent

PAGES = [
    [
        "Introduction",
        "The lighthouse keepers' guild grew by a third this year, e.g. in the",
        "northern stations. Its newsletter carried more infor-",
        "mation than ever before.",
    ],
    [
        "Outlook",
        "Dr. Okafor expects the U.S. stations to follow. Write to",
        "guild@example.org or see https://example.org/reports/2025 for more.",
    ],
]


def pdf() -> None:
    canvas = Canvas(str(HERE / "report.pdf"), pagesize=A5)
    for number, lines in enumerate(PAGES, 1):
        canvas.setFont("Helvetica", 9)
        canvas.drawString(40, 560, "Lighthouse Guild Annual Report 2025")
        canvas.setFont("Helvetica", 11)
        y = 520
        for line in lines:
            canvas.drawString(40, y, line)
            y -= 16
        canvas.setFont("Helvetica", 9)
        canvas.drawString(40, 30, f"Page {number} of {len(PAGES)}")
        canvas.showPage()
    canvas.save()


def docx() -> None:
    doc = Document()
    doc.add_heading("RESULTS BY STATION", level=1)
    doc.add_paragraph("Each station logged its clear nights, i.e. nights with no fog.")
    table = doc.add_table(rows=3, cols=2)
    for row, cells in zip(
        table.rows,
        [("Station", "Clear nights"), ("North Point", "212"), ("Gull Rock", "187")],
        strict=True,
    ):
        for cell, text in zip(row.cells, cells, strict=True):
            cell.text = text
    doc.add_heading("Next steps", level=2)
    for item in [
        "Repaint the north tower",
        "Order lamp oil vs. buying it locally",
        "Contact keeper@example.org",
    ]:
        doc.add_paragraph(item, style="List Bullet")
    doc.save(str(HERE / "stations.docx"))


def pptx() -> None:
    deck = Presentation()
    for title, points in [
        ("Guild Meeting", ["Autumn 2025", "Harbour hall, 7 pm"]),
        (
            "Agenda",
            ["Fog horn schedule", "New keepers, e.g. two from Gull Rock", "Any other business"],
        ),
    ]:
        slide = deck.slides.add_slide(deck.slide_layouts[1])
        slide.shapes.title.text = title
        body = slide.placeholders[1].text_frame
        body.text = points[0]
        for point in points[1:]:
            body.add_paragraph().text = point
    deck.save(str(HERE / "meeting.pptx"))


if __name__ == "__main__":
    pdf()
    docx()
    pptx()
